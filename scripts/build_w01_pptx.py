"""1주차 「여러 가지 순열」 PPT 생성.

먼저 render_w01.py 로 애니메이션을 렌더링한 뒤 실행.
"""
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
IMG = ROOT / "04_이미지" / "01주차"
ANIM = ROOT / "05_결과물" / "01주차" / "애니메이션"
OUT = ROOT / "05_결과물" / "01주차" / "01주차_여러가지순열.pptx"

FONT = "맑은 고딕"
MATH = "Cambria Math"
C = {k: RGBColor.from_string(v) for k, v in {
    "ink": "1D2433", "sub": "5B6475", "line": "D5DAE3", "white": "FFFFFF", "panel": "F4F6FA",
    "blue": "2F6FDE", "blue_soft": "DCE8FB", "red": "E0533D", "red_soft": "FBE1DC",
    "gold": "F2B632", "gold_soft": "FDF0CF", "cover_sub": "C6CEDE", "cover_acc": "9FB8E8",
}.items()}

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
page_no = 0


# ---------- 기본 도구 ----------
def seg_runs(p, segs, size, color, bold, font):
    """segs: 문자열 또는 (텍스트, {color, bold, size, font, sub, sup}) 목록"""
    if isinstance(segs, str):
        segs = [(segs, {})]
    for s in segs:
        txt, st = (s, {}) if isinstance(s, str) else s
        r = p.add_run()
        r.text = txt
        f = r.font
        f.size = Pt(st.get("size", size))
        f.bold = st.get("bold", bold)
        f.color.rgb = C[st.get("color", color)]
        f.name = st.get("font", font)
        ea = r._r.get_or_add_rPr()
        ea_el = ea.find(qn("a:ea"))
        if ea_el is None:
            ea_el = ea.makeelement(qn("a:ea"), {})
            ea.append(ea_el)
        ea_el.set("typeface", st.get("font", font))
        if st.get("sub"):
            ea.set("baseline", "-25000")
        if st.get("sup"):
            ea.set("baseline", "30000")


def text(slide, x, y, w, h, paras, size=20, color="ink", bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT, spacing=1.15, space_after=0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    paras = paras if isinstance(paras, list) else [paras]
    for i, segs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p._p.get_or_add_pPr().set("eaLnBrk", "0")
        p.line_spacing = spacing
        p.space_after = Pt(space_after)
        seg_runs(p, segs, size, color, bold, font)
    return tb


def box(slide, x, y, w, h, fill="white", line="line", lw=1.5, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = C[fill]
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = C[line]
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def label(slide, x, y, w, h, segs, fill="white", line="line", size=20, color="ink", bold=True, lw=1.5,
          shape=MSO_SHAPE.ROUNDED_RECTANGLE, font=FONT):
    s = box(slide, x, y, w, h, fill, line, lw, shape)
    tf = s.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p._p.get_or_add_pPr().set("eaLnBrk", "0")
    seg_runs(p, segs, size, color, bold, font)
    return s


def arrow(slide, x1, y1, x2, y2, color="blue", w=2.25):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = C[color]
    ln.line.width = Pt(w)
    tail = ln.line._get_or_add_ln().makeelement(qn("a:tailEnd"), {"type": "triangle"})
    ln.line._get_or_add_ln().append(tail)
    return ln


def photo(slide, path, x, y, w, h, caption=None):
    """사진을 상자 비율에 맞게 가운데 기준으로 잘라 넣기"""
    iw, ih = Image.open(path).size
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    target, actual = w / h, iw / ih
    if actual > target:
        cut = (1 - target / actual) / 2
        pic.crop_left = pic.crop_right = cut
    else:
        cut = (1 - actual / target) / 2
        pic.crop_top = pic.crop_bottom = cut
    pic.line.color.rgb = C["line"]
    pic.line.width = Pt(1)
    if caption:
        text(slide, x, y + h + 0.08, w, 0.8, caption, size=15, color="sub", spacing=1.1)
    return pic


def movie(slide, name, x, y, w):
    h = w * 9 / 16
    slide.shapes.add_movie(str(ANIM / f"{name}.mp4"), Inches(x), Inches(y), Inches(w), Inches(h),
                           poster_frame_image=str(ANIM / f"{name}_마지막장면.png"), mime_type="video/mp4")
    box(slide, x, y, w, h, fill=None, line="line", lw=1)
    text(slide, x, y + h + 0.03, w, 0.3, "▶ 화면을 클릭하면 애니메이션 재생", size=12, color="sub", align=PP_ALIGN.RIGHT)


def bullets(slide, x, y, w, h, items, size=22, gap=10):
    paras = [[("■  ", {"color": "blue", "size": size - 8})] + (it if isinstance(it, list) else [it]) for it in items]
    return text(slide, x, y, w, h, paras, size=size, spacing=1.25, space_after=gap)


def new_slide(title=None, sub=None, src=None, notes=None):
    global page_no
    page_no += 1
    s = prs.slides.add_slide(BLANK)
    if title:
        segs = [(title, {"bold": True, "size": 32})]
        if sub:
            segs.append(("   " + sub, {"size": 16, "color": "sub"}))
        text(s, 0.6, 0.38, 12.1, 0.7, [segs], anchor=MSO_ANCHOR.BOTTOM)
        box(s, 0.6, 1.12, 12.13, 0.035, fill="ink", line=None, shape=MSO_SHAPE.RECTANGLE)
        text(s, 12.2, 7.05, 0.8, 0.3, str(page_no), size=12, color="sub", align=PP_ALIGN.RIGHT)
    if src:
        text(s, 0.6, 7.05, 8, 0.3, "근거: " + src, size=12, color="sub")
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


K = lambda t: (t, {"color": "blue", "bold": True})
R = lambda t: (t, {"color": "red", "bold": True})
SUB = lambda t: (t, {"sub": True})
SUP = lambda t: (t, {"sup": True})


# ---------- 슬라이드 ----------
# 1. 표지
s = new_slide(notes="1주차 도입. 지난 학기 순열(nPr)과 계승(n!)을 복습한 뒤 시작.")
bg = s.background.fill
bg.solid()
bg.fore_color.rgb = C["ink"]
text(s, 0.9, 2.2, 11, 0.5, "확률과 통계 · 1주차", size=22, color="cover_acc", bold=True)
text(s, 0.9, 2.75, 11, 1.3, "여러 가지 순열", size=60, color="white", bold=True)
text(s, 0.9, 4.1, 11, 0.6, "원순열 · 중복순열 · 같은 것이 있는 순열", size=24, color="cover_sub")
text(s, 0.9, 6.7, 11, 0.4, "근거: 미래엔 『확률과 통계』 Ⅰ. 경우의 수 – 1. 순열과 조합 (pp.10–16)", size=13, color="cover_sub")

# 2. 학습 흐름
s = new_slide("오늘의 학습 흐름", notes="세 가지 순열을 일상 비유와 함께 미리 소개.")
cards = [
    ("원순열", "원탁_윈체스터.jpg", "원형으로 배열,\n회전하여 같은 배열은 1가지", "비유: 상석이 없는 아서왕의 원탁"),
    ("중복순열", "번호자물쇠.jpg", "같은 것을 여러 번\n뽑을 수 있는 배열", "비유: 숫자를 다시 써도 되는 자물쇠"),
    ("같은 것이 있는 순열", "모스전신기.jpg", "구별되지 않는 것이\n섞여 있는 배열", "비유: 점과 선만 반복하는 모스 부호"),
]
for i, (name, img, desc, ana) in enumerate(cards):
    x = 0.6 + i * 4.1
    box(s, x, 1.5, 3.85, 5.35)
    label(s, x + 1.67, 1.68, 0.5, 0.5, str(i + 1), fill="blue", line=None, size=16, color="white", shape=MSO_SHAPE.OVAL)
    photo(s, IMG / img, x + 0.25, 2.35, 3.35, 2.25)
    text(s, x + 0.2, 4.75, 3.45, 0.5, name, size=24, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 5.3, 3.45, 0.8, desc, size=16, color="sub", align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 6.1, 3.45, 0.5, ana, size=15, color="blue", bold=True, align=PP_ALIGN.CENTER)

# 3. 원순열 도입
s = new_slide("원순열", "회전하면 같은 배열", "미래엔 p.11",
              notes="애니메이션: 원탁을 90°씩 돌리며 ABCD, DABC, CDAB, BCDA가 같은 배열임을 확인.\n질문: 원탁에는 왜 '첫 번째 자리'가 없을까?")
photo(s, IMG / "원탁_윈체스터.jpg", 0.6, 1.5, 4.4, 3.3,
      [[("비유: 아서왕의 원탁", {"bold": True, "color": "ink"})], "상석이 없어 어느 자리든 동등한 배치"])
movie(s, "CircleRotate", 5.35, 1.5, 7.4)
bullets(s, 0.6, 5.95, 12.1, 1.0, [
    ["일렬 배열에서는 서로 다른 4가지, 원형 배열에서는 ", K("모두 같은 1가지")],
    "기준이 되는 “첫 자리”가 없는 것이 원형 배열의 특징",
], size=20, gap=4)

# 4. 원순열의 수
s = new_slide("원순열의 수", src="미래엔 pp.11–12",
              notes="일렬 배열 24가지를 회전 묶음(한 줄 4가지)으로 묶으면 6묶음. 일반화하면 n!/n = (n−1)!.")
movie(s, "CircleCount", 0.6, 1.45, 7.6)
box(s, 8.5, 1.45, 4.25, 3.4, fill="blue_soft", line=None)
text(s, 8.65, 1.6, 3.95, 0.4, "서로 다른 n개의 원순열의 수", size=17, color="blue", bold=True, align=PP_ALIGN.CENTER)
text(s, 8.65, 2.1, 3.95, 0.8, "n! ÷ n = (n − 1)!", size=32, bold=True, font=MATH, align=PP_ALIGN.CENTER)
text(s, 8.75, 3.05, 3.85, 1.7, [
    [("n!", {"font": MATH, "bold": True}), "   n개를 일렬로 배열하는 경우의 수"],
    [("÷ n", {"font": MATH, "bold": True}), "   겹치는 n가지를 1가지로 묶음"],
], size=15, spacing=1.2, space_after=8)
bullets(s, 8.5, 5.15, 4.3, 1.6, [
    ["4개의 원순열", K(" 24 ÷ 4 = 3! = 6")],
    ["5개의 원순열", K(" 4! = 24")],
], size=19, gap=6)

# 5. 원순열 예제
s = new_slide("예제: 이웃하여 앉는 경우", "묶어서 하나로 취급", "미래엔 p.12 예제 1",
              notes="이웃하는 조건 → 한 묶음으로 보고 원순열, 마지막에 묶음 내부 자리 바꿈을 곱함.")
movie(s, "Adjacent", 0.6, 1.45, 7.6)
label(s, 8.5, 1.5, 4.25, 0.5, "회장, 부회장 포함 5명의 원탁 회의", fill="blue_soft", line=None, size=15, color="blue")
bullets(s, 8.5, 2.25, 4.3, 4.5, [
    ["5명이 앉는 경우 ", K("(5−1)! = 24")],
    "회장·부회장 묶음과 3명의 원순열 (4−1)! = 6",
    "묶음 안 자리 바꿈 2! = 2",
    ["이웃하는 경우 ", K("6 × 2 = 12")],
], size=19, gap=10)

# 6. 중복순열 도입
s = new_slide("중복순열", "같은 것을 다시 뽑을 수 있는 배열", "미래엔 p.13",
              notes="자물쇠 사진으로 시작: 같은 숫자를 다시 돌려 쓸 수 있음. 애니메이션 후반은 비트가 하나 늘 때마다 경우가 2배가 되는 트리.")
photo(s, IMG / "번호자물쇠.jpg", 0.6, 1.5, 4.4, 3.3,
      [[("비유: 여행가방 번호 자물쇠", {"bold": True, "color": "ink"})], "자리마다 0~9 중 아무 숫자나 다시 사용 가능"])
movie(s, "Repeated", 5.35, 1.5, 7.4)
bullets(s, 0.6, 5.95, 12.1, 1.0, [
    ["세 자리 비밀번호 10 × 10 × 10 = ", K("10³ = 1000가지")],
    ["8비트 신호 2⁸ = ", K("256가지"), ": 컴퓨터 1바이트가 나타내는 값의 개수"],
], size=20, gap=4)

# 7. 중복순열의 수
s = new_slide("중복순열의 수", src="미래엔 pp.13–14",
              notes="기호 Π는 Product의 P. 순열 nPr은 r ≤ n이어야 하지만 중복순열은 r > n도 가능.")
box(s, 0.6, 1.5, 6.0, 4.0, fill="blue_soft", line=None)
text(s, 0.8, 1.65, 5.6, 0.7, "서로 다른 n개에서 중복을 허용하여 r개를 택하는 중복순열의 수",
     size=17, color="blue", bold=True, align=PP_ALIGN.CENTER)
text(s, 0.8, 2.4, 5.6, 0.9, [[SUB("n"), "Π", SUB("r"), " = n × n × ⋯ × n = n", SUP("r")]],
     size=32, bold=True, font=MATH, align=PP_ALIGN.CENTER)
text(s, 1.0, 3.45, 5.3, 2.0, [
    [("Π", {"font": MATH, "bold": True}), "   Product(곱)의 P에 해당하는 그리스 문자"],
    [("n", {"font": MATH, "bold": True}), "   각 자리에 올 수 있는 후보의 개수"],
    [("r", {"font": MATH, "bold": True}), "   자리의 개수 (n보다 커도 가능)"],
], size=16, spacing=1.2, space_after=8)
slots = ["첫 번째", "두 번째", "세 번째", "⋯", "r 번째"]
for i, sl in enumerate(slots):
    x = 7.1 + i * 1.12
    if sl == "⋯":
        text(s, x, 1.75, 0.95, 0.6, "⋯", size=26, color="sub", align=PP_ALIGN.CENTER)
        continue
    label(s, x, 1.6, 0.95, 0.75, sl, fill="blue_soft", line="blue", size=14, color="blue")
    arrow(s, x + 0.475, 3.05, x + 0.475, 2.45)
    text(s, x - 0.1, 3.1, 1.15, 0.45, "n가지", size=17, bold=True, align=PP_ALIGN.CENTER)
text(s, 7.1, 3.6, 5.6, 0.4, "곱의 법칙: n을 r번 곱하기", size=15, color="sub", align=PP_ALIGN.CENTER)
bullets(s, 7.1, 4.3, 5.7, 2.4, [
    [SUB("5"), "Π", SUB("2"), " = 5² = 25,   ", SUB("3"), "Π", SUB("4"), " = 3⁴ = 81"],
    ["순열 ", SUB("n"), "P", SUB("r"), "과 달리 ", K("r > n"), "도 가능한 점"],
], size=20, gap=10)

# 8. 중복순열 예제
s = new_slide("예제: 맨 앞자리 조건이 있는 경우", src="미래엔 p.14 예제 2",
              notes="천의 자리에 0이 올 수 없다는 조건을 먼저 처리하고 나머지 세 자리는 중복순열.")
places = [("천의 자리", "1, 2, 3\n(0 제외)", "3가지", True), ("백의 자리", "0, 1, 2, 3", "4가지", False),
          ("십의 자리", "0, 1, 2, 3", "4가지", False), ("일의 자리", "0, 1, 2, 3", "4가지", False)]
for i, (nm, cand, cnt, first) in enumerate(places):
    x = 0.7 + i * 1.85
    text(s, x, 1.6, 1.6, 0.4, nm, size=16, color="sub", bold=True, align=PP_ALIGN.CENTER)
    label(s, x, 2.05, 1.6, 2.0, cand, fill="red_soft" if first else "blue_soft", line="red" if first else "blue",
          size=18, color="red" if first else "blue", lw=2.25)
    text(s, x, 4.15, 1.6, 0.5, cnt, size=22, bold=True, align=PP_ALIGN.CENTER)
box(s, 2.55, 4.8, 5.4, 0.04, fill="blue", line=None, shape=MSO_SHAPE.RECTANGLE)
text(s, 2.55, 4.9, 5.4, 0.5, [[SUB("4"), "Π", SUB("3"), " = 4³ = 64"]], size=20, color="blue", bold=True, align=PP_ALIGN.CENTER)
label(s, 8.4, 1.6, 4.35, 0.5, "숫자 0, 1, 2, 3으로 만드는 네 자리 자연수", fill="blue_soft", line=None, size=15, color="blue")
bullets(s, 8.4, 2.35, 4.4, 4.0, [
    ["천의 자리: 0 제외, ", R("3가지")],
    ["나머지 세 자리: ", SUB("4"), "Π", SUB("3"), " = 64"],
    ["전체 3 × 64 = ", K("192개")],
], size=20, gap=12)

# 9. 같은 것이 있는 순열 도입
s = new_slide("같은 것이 있는 순열", "구별되지 않는 것이 섞인 배열", "미래엔 p.15",
              notes="애니메이션: a₁a₂a₃b₁b₂를 모두 다르게 보면 120가지, 첨자를 지우면 3!×2! = 12가지씩 겹침 → 10가지.")
photo(s, IMG / "모스전신기.jpg", 0.6, 1.5, 4.4, 3.3,
      [[("비유: 모스 전신기", {"bold": True, "color": "ink"})], "점(·)과 선(−) 두 기호만 반복해 문자를 표현"])
movie(s, "SameThings", 5.35, 1.5, 7.4)
bullets(s, 0.6, 5.95, 12.1, 1.0, [
    "모두 다르다고 보고 세면 같은 배열이 여러 번 중복",
    ["같은 것끼리의 자리 바꿈 수 ", K("3! × 2!"), "로 나누어 중복 제거"],
], size=20, gap=4)

# 10. 같은 것이 있는 순열의 수 + 모스 부호
s = new_slide("같은 것이 있는 순열의 수", src="미래엔 pp.14–15",
              notes="교과서 생각 열기의 모스 부호 예: 점 1개, 선 2개를 전부 나열 → 3!/2! = 3가지.")
box(s, 0.6, 1.5, 6.0, 3.6, fill="blue_soft", line=None)
text(s, 0.8, 1.65, 5.6, 0.7, "같은 것이 각각 p개, q개, …, r개 있는 n개를 일렬로 나열하는 순열의 수",
     size=17, color="blue", bold=True, align=PP_ALIGN.CENTER)
text(s, 0.8, 2.4, 5.6, 0.8, "n! ÷ (p! × q! × ⋯ × r!)", size=30, bold=True, font=MATH, align=PP_ALIGN.CENTER)
text(s, 1.0, 3.4, 5.3, 1.6, [
    [("n!", {"font": MATH, "bold": True}), "   모두 다르다고 볼 때의 배열 수"],
    [("p!, q!", {"font": MATH, "bold": True}), "   같은 것끼리 자리 바꿈으로 생기는 중복"],
    [("조건", {"bold": True}), "   p + q + ⋯ + r = n"],
], size=16, spacing=1.2, space_after=6)
bullets(s, 0.6, 5.4, 6.0, 1.4, [["a 3개, b 2개: 5! ÷ (3! × 2!) = ", K("10")]], size=20)
label(s, 7.1, 1.5, 5.65, 0.5, "모스 부호: 점 1개, 선 2개를 전부 나열", fill="blue_soft", line=None, size=15, color="blue")
for k, code in enumerate(["·−−", "−·−", "−−·"]):
    y = 2.3 + k * 0.95
    box(s, 7.1, y, 5.65, 0.75, fill="panel", line=None)
    for j, ch in enumerate(code):
        x = 7.6 + j * 1.1
        if ch == "·":
            box(s, x + 0.3, y + 0.24, 0.27, 0.27, fill="red", line=None, shape=MSO_SHAPE.OVAL)
        else:
            box(s, x, y + 0.29, 0.9, 0.17, fill="ink", line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 11.2, y + 0.15, 1.4, 0.45, f"{k + 1}번째", size=15, color="sub", align=PP_ALIGN.RIGHT)
text(s, 7.1, 5.25, 5.65, 0.6, [["모스 부호의 수  3! ÷ 2! = ", K("3가지")]], size=22, bold=True, align=PP_ALIGN.CENTER)
text(s, 7.1, 5.95, 5.65, 0.8, "생각해 보기: 점 2개, 선 2개를 전부 나열하면 몇 가지?", size=16, color="sub", align=PP_ALIGN.CENTER)

# 11. 최단 경로
s = new_slide("예제: 격자에서의 최단 경로", "같은 것이 있는 순열의 활용", "미래엔 p.16 예제 3",
              notes="애니메이션 후반: 각 교차점까지의 경로 수 = 왼쪽 + 아래쪽 (파스칼 삼각형과 같은 구조). B에서 15.")
photo(s, IMG / "격자도시_바르셀로나.jpg", 0.6, 1.5, 4.4, 3.3,
      [[("비유: 바르셀로나 격자 도시", {"bold": True, "color": "ink"})], "동쪽·북쪽으로만 가는 최단 이동", "회로 배선·로봇 경로와 같은 구조"])
movie(s, "ShortestPath", 5.35, 1.5, 7.4)
bullets(s, 0.6, 5.95, 12.1, 1.0, [
    ["오른쪽 한 칸 ", K("a"), ", 위쪽 한 칸 ", R("b"), " → 최단 경로 = a 4개, b 2개의 나열"],
    ["경로의 수 6! ÷ (4! × 2!) = ", K("15")],
], size=20, gap=4)

# 12. 요약
s = new_slide("요약", notes="세 공식의 공통점: 모두 '일단 다 다르게 세고, 겹친 만큼 나누거나 곱하는' 구조.")
summ = [
    ("원순열", [["(n − 1)!"]], "회전하여 같은 n가지를\n1가지로 묶는 계산", "원탁_윈체스터.jpg"),
    ("중복순열", [[SUB("n"), "Π", SUB("r"), " = n", SUP("r")]], "자리마다 n가지,\nr번 곱하는 계산", "번호자물쇠.jpg"),
    ("같은 것이 있는 순열", [["n! ÷ (p! q! ⋯ r!)"]], "같은 것끼리의 자리 바꿈\n중복을 나누는 계산", "모스전신기.jpg"),
]
for i, (nm, f, desc, img) in enumerate(summ):
    x = 0.6 + i * 4.1
    box(s, x, 1.5, 3.85, 5.35)
    photo(s, IMG / img, x + 0.25, 1.75, 3.35, 1.9)
    text(s, x + 0.2, 3.85, 3.45, 0.5, nm, size=22, bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 4.45, 3.45, 0.7, f, size=26, bold=True, color="blue", font=MATH, align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 5.4, 3.45, 1.0, desc, size=16, color="sub", align=PP_ALIGN.CENTER)

# 13. 연습문제
s = new_slide("연습문제", notes="정답은 다음 슬라이드. 3번은 같은 것이 있는 순열과 최단 경로(예제 3)가 같은 답임에 주목.")
ex = [
    ("서버랙.jpg", "서버 6대를 원형 링 네트워크로\n연결하는 배치의 수", "(회전하여 같은 것은 같은 배치)"),
    ("RGB_LED.jpg", "4가지 색을 낼 수 있는 LED 3개로\n표시 가능한 신호의 수", "(같은 색 중복 가능)"),
    (None, "1 네 개, 0 두 개로 만드는\n6비트 신호의 수", ""),
]
for i, (img, q, cond) in enumerate(ex):
    x = 0.6 + i * 4.1
    box(s, x, 1.5, 3.85, 5.35)
    label(s, x + 1.67, 1.68, 0.5, 0.5, str(i + 1), fill="blue", line=None, size=16, color="white", shape=MSO_SHAPE.OVAL)
    if img:
        photo(s, IMG / img, x + 0.25, 2.35, 3.35, 2.25)
    else:
        for j, b in enumerate("101101"):
            label(s, x + 0.33 + j * 0.53, 3.05, 0.45, 0.75, b, fill="blue_soft" if b == "1" else "white",
                  line="blue" if b == "1" else "ink", size=20, color="blue" if b == "1" else "ink")
    text(s, x + 0.1, 4.8, 3.65, 1.1, q, size=17, align=PP_ALIGN.CENTER)
    text(s, x + 0.2, 5.95, 3.45, 0.5, cond, size=14, color="sub", align=PP_ALIGN.CENTER)

# 14. 정답
s = new_slide("연습문제 정답")
bullets(s, 0.6, 1.6, 12.1, 4.5, [
    [("1번", {"bold": True}), "  6개의 원순열 (6 − 1)! = 5! = ", K("120")],
    [("2번", {"bold": True}), "  4가지 색에서 3개를 택하는 중복순열 ", SUB("4"), "Π", SUB("3"), " = 4³ = ", K("64")],
    [("3번", {"bold": True}), "  같은 것이 있는 순열 6! ÷ (4! × 2!) = ", K("15"), "  (격자 최단 경로와 같은 구조)"],
], size=24, gap=18)

# 15. 이미지 출처
s = new_slide("이미지 출처", notes="모든 사진은 Wikimedia Commons의 자유 라이선스 이미지. CC BY-SA 이미지는 같은 조건으로만 재배포 가능.")
credits = {c["file"]: c for c in json.load(open(IMG / "출처.json", encoding="utf-8"))}
used = ["원탁_윈체스터.jpg", "번호자물쇠.jpg", "모스전신기.jpg", "격자도시_바르셀로나.jpg", "서버랙.jpg", "RGB_LED.jpg"]
paras = []
for f in used:
    c = credits[f]
    paras.append([(f.rsplit(".", 1)[0].replace("_", " "), {"bold": True}),
                  f"  {c['title'].removeprefix('File:')} · {c['author']} · {c['license']}"])
    paras.append([(c["source"], {"color": "sub", "size": 11})])
text(s, 0.6, 1.5, 12.1, 5.4, paras, size=14, spacing=1.1, space_after=4)
text(s, 0.6, 6.6, 12.1, 0.4, "출처: Wikimedia Commons", size=12, color="sub")

prs.save(OUT)
print("saved", OUT, page_no, "slides")
